import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from backend.services.skills import find_skills

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt"}

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Indian mobile: optional +91, then 10 digits starting 6-9, optionally split 5+5.
_PHONE = re.compile(r"(?<!\d)(?:\+?91[\s-]*)?([6-9]\d{4})[\s-]?(\d{5})(?!\d)")
_EXPERIENCE = re.compile(
    r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\b[^.\n]{0,30}?\bexperience", re.IGNORECASE
)
_FRESHER = re.compile(r"\bfresher\b", re.IGNORECASE)
_DEGREE = re.compile(
    r"\b(?:b\.?\s?tech|m\.?\s?tech|b\.e\.?|m\.e\.?|b\.?sc|m\.?sc|bca|mca|bba|mba|b\.?com|m\.?com"
    r"|bachelor(?:'?s)?|master(?:'?s)?\s+(?:of|in)|diploma|ph\.?d)(?![a-z])",
    re.IGNORECASE,
)
_NAME_WORD = re.compile(r"[A-Za-z][A-Za-z.'-]*")
_NOT_A_NAME = {"resume", "curriculum", "vitae", "cv", "profile", "biodata", "bio-data"}


class CVReadError(Exception):
    """The file could not be read as a CV. The message is safe to show to the user."""


@dataclass
class ParsedCV:
    full_name: str = ""
    email: str = ""
    phone: str = ""
    experience_years: float | None = None
    skills: list[str] = field(default_factory=list)
    education: str = ""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise CVReadError("Only PDF, DOCX and TXT files are supported.")
    try:
        if suffix == ".txt":
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".pdf":
            return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        return _docx_text(path)
    except Exception as exc:  # pypdf and python-docx raise many unrelated types on bad files
        raise CVReadError("This file could not be opened. It may be damaged or password-protected.") from exc


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # CV templates often put everything inside tables, which .paragraphs does not include.
    for table in document.tables:
        for row in table.rows:
            parts.append("  ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def parse_cv(text: str) -> ParsedCV:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    email = _EMAIL.search(text)
    phone = _PHONE.search(text)
    return ParsedCV(
        full_name=_guess_name(lines),
        email=email.group(0) if email else "",
        phone="".join(phone.groups()) if phone else "",
        experience_years=_guess_experience(text),
        skills=find_skills(text),
        education=_guess_education(lines),
    )


def _guess_name(lines: list[str]) -> str:
    for line in lines[:5]:
        words = line.split()
        if not 2 <= len(words) <= 4:
            continue
        if not all(_NAME_WORD.fullmatch(word) for word in words):
            continue
        if any(word.lower() in _NOT_A_NAME for word in words):
            continue
        return line.title() if line.isupper() else line
    return ""


def _guess_experience(text: str) -> float | None:
    match = _EXPERIENCE.search(text)
    if match:
        return float(match.group(1))
    if _FRESHER.search(text):
        return 0.0
    return None


def _guess_education(lines: list[str]) -> str:
    found: list[str] = []
    for line in lines:
        if _DEGREE.search(line) and line not in found:
            found.append(line[:200])
        if len(found) == 3:
            break
    return "\n".join(found)
