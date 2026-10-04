import pytest
from docx import Document

from backend.services.cv_parser import CVReadError, extract_text, parse_cv
from backend.services.skills import canonical_skill, find_skills

SAMPLE = """RAVI KUMAR
Frontend Developer
ravi.kumar99@gmail.com | +91 98765 43210 | Hyderabad

SUMMARY
Fresher with strong basics in web development.

SKILLS
JavaScript, React.js, Next.js, HTML5, CSS3, Tailwind CSS, Node.js, Git, C++

EDUCATION
B.Tech in Computer Science, JNTU Hyderabad, 2026
"""


def test_parse_cv_reads_the_basics():
    parsed = parse_cv(SAMPLE)
    assert parsed.full_name == "Ravi Kumar"
    assert parsed.email == "ravi.kumar99@gmail.com"
    assert parsed.phone == "9876543210"
    assert parsed.experience_years == 0.0
    assert "B.Tech in Computer Science" in parsed.education
    for skill in ["JavaScript", "React", "Next.js", "HTML", "CSS", "Tailwind CSS", "Node.js", "Git", "C++"]:
        assert skill in parsed.skills


def test_parse_cv_reads_years_of_experience():
    assert parse_cv("I have 2.5 years of professional experience in testing").experience_years == 2.5
    assert parse_cv("3+ yrs experience").experience_years == 3.0
    assert parse_cv("No mention at all").experience_years is None


def test_skill_matching_respects_word_edges():
    assert find_skills("Worked with JavaScript daily") == ["JavaScript"]  # not Java
    assert "Git" not in find_skills("Code is on GitHub")
    assert "SQL" not in find_skills("Used MySQL and PostgreSQL")
    assert "JavaScript" not in find_skills("Built APIs in Node.js")  # "js" inside node.js
    assert find_skills("ASP.NET with C#") == ["C#", "ASP.NET"]


def test_canonical_skill_merges_spellings():
    assert canonical_skill("reactjs") == "React"
    assert canonical_skill("  postgres ") == "PostgreSQL"
    assert canonical_skill("Some Niche Tool") == "Some Niche Tool"


def test_name_is_not_guessed_from_a_heading():
    assert parse_cv("CURRICULUM VITAE\nemail: a@b.com").full_name == ""


def test_extract_text_from_docx_includes_tables(tmp_path):
    path = tmp_path / "cv.docx"
    document = Document()
    document.add_paragraph("Asha Verma")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, Django"
    document.save(str(path))

    text = extract_text(path)
    assert "Asha Verma" in text
    assert "Python, Django" in text


def test_extract_text_from_pdf(tmp_path, make_pdf):
    path = tmp_path / "cv.pdf"
    path.write_bytes(make_pdf("Priya Sharma knows Python"))
    assert "Priya Sharma knows Python" in extract_text(path)


def test_damaged_file_gives_a_readable_error(tmp_path):
    path = tmp_path / "cv.pdf"
    path.write_bytes(b"this is not a pdf")
    with pytest.raises(CVReadError):
        extract_text(path)
